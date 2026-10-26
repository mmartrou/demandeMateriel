"""Écriture des fichiers Excel du planning (planning optimisé ou planning sauvegardé)."""
import database
import site_settings

from .cours import extract_material_needs, get_room_layout, h_to_min
from .impression import rafraichir_cours


def _ecrire_entetes_de_groupes(ws, groupes, salle_list, font, border):
    """Ligne 1 de la feuille : un encadré fusionné par groupe de salles (ex : Physique, Chimie)."""
    from openpyxl.styles import Alignment
    for libelle, membres in groupes:
        colonnes = [2 + salle_list.index(s) for s in membres if s in salle_list]
        if not colonnes:
            continue
        if len(colonnes) > 1 and colonnes[-1] > colonnes[0]:
            ws.merge_cells(start_row=1, start_column=colonnes[0], end_row=1, end_column=colonnes[-1])
        for col in range(colonnes[0], colonnes[-1] + 1):
            cell = ws.cell(row=1, column=col, value=libelle if col == colonnes[0] else None)
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border


def generer_excel_optimise(cours, salles, x, solver, unassigned_courses, date_param=None, custom_room_assignments=None):
    """Génération Excel optimisée avec le solveur CP - Style grille horaire."""
    try:
        print(f"🔧 [Excel] Début génération Excel avec custom_room_assignments type: {type(custom_room_assignments)}")
        print(f"🔧 [Excel] custom_room_assignments value: {custom_room_assignments}")
        
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
        from openpyxl.worksheet.page import PageMargins
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont

        wb = Workbook()

        # Créneaux horaires personnalisés (comme dans main.py)
        layout_horaires = site_settings.get_time_layout()
        horaires_str = layout_horaires['grid_rows']
        separateurs = set(layout_horaires['separators'])
        horaires = [h_to_min(h) for h in horaires_str]

        JOURS_FR = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche']

        def _row_border(idx_h):
            """Bordure horizontale alternée (pleine/pointillés), avec séparateurs
            renforcés aux lignes de pause réglées dans /admin/establishment."""
            is_major = horaires_str[idx_h] in separateurs
            next_is_major = idx_h + 1 < len(horaires_str) and horaires_str[idx_h + 1] in separateurs
            top_style = 'thick' if is_major else ('medium' if idx_h % 2 == 0 else 'dashed')
            bottom_style = 'thick' if next_is_major else ('medium' if idx_h % 2 == 0 else 'dashed')
            return Border(left=Side(style='thick'), right=Side(style='thick'),
                          top=Side(style=top_style), bottom=Side(style=bottom_style))

        def _teacher_rich_text(teacher, rest_lines):
            """Nom du professeur en gras, 1.5x plus grand que le reste du contenu."""
            rest = "\n".join(l for l in rest_lines if l)
            runs = [TextBlock(InlineFont(rFont="Calibri", b=True, sz=17), teacher or "")]
            if rest:
                runs.append(TextBlock(InlineFont(rFont="Calibri", sz=11), "\n" + rest))
            return CellRichText(*runs)

        def _apply_print_setup(ws, n_rooms, n_rows):
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3, header=0, footer=0)
            last_col_letter = ws.cell(row=2, column=1 + n_rooms).column_letter
            ws.print_area = f"A1:{last_col_letter}{2 + n_rows}"
        
        # Ordre des salles et regroupements d'en-tête : réglés dans /admin/rooms
        ordered_rooms, groupes_salles = get_room_layout()

        # Filtrer pour ne garder que les salles qui existent
        salle_list = [room for room in ordered_rooms if room in salles.keys()]
        
        # Ajouter les salles restantes qui ne sont pas dans l'ordre personnalisé (au cas où)
        remaining_rooms = [room for room in salles.keys() if room not in salle_list]
        salle_list.extend(sorted(remaining_rooms))
        
        # Feuille 1: Planning détaillé (techniciens)
        ws1 = wb.active
        ws1.title = "Planning_Techniciens"
        
        # Styles
        bold_font = Font(bold=True, size=16)
        header_font = Font(bold=True, size=36)
        border = Border(left=Side(style='thick'), right=Side(style='thick'), 
                       top=Side(style='thick'), bottom=Side(style='thick'))
        
        # Date / jour de la semaine - Utiliser la date du planning, pas la date actuelle
        from datetime import datetime
        if date_param:
            if isinstance(date_param, str):
                planning_date = datetime.strptime(date_param, '%Y-%m-%d')
            else:
                planning_date = date_param
            formatted_date = planning_date.strftime('%d/%m')
        else:
            planning_date = datetime.now()
            formatted_date = planning_date.strftime('%d/%m')
        jour_semaine = JOURS_FR[planning_date.weekday()]

        # En-tête principal (ligne 1)
        ws1.cell(row=1, column=1, value=jour_semaine)
        ws1.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws1.column_dimensions['A'].width = 12
        ws1.row_dimensions[1].height = 40
        ws1.row_dimensions[2].height = 40

        # Encadrés de groupes de salles (Physique, Chimie, ...)
        _ecrire_entetes_de_groupes(ws1, groupes_salles, salle_list, header_font, border)

        # Noms des salles (ligne 2)
        for idx_s, s in enumerate(salle_list):
            col_letter = ws1.cell(row=2, column=2+idx_s).column_letter
            cell = ws1.cell(row=2, column=2+idx_s, value=s)
            ws1.column_dimensions[col_letter].width = 15
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        # Date (ligne 2, colonne 1)
        ws1.cell(row=2, column=1, value=formatted_date)
        ws1.cell(row=2, column=1).font = bold_font
        ws1.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=2, column=1).border = border
        
        # Préparation matrice pour fusion de cellules
        cell_matrix = [[{"content": "", "content_techniciens": "", "content_affichage": "", "merge": False, "merge_len": 1, "matiere": ""} for _ in salle_list] for _ in horaires]
        
        # Remplir la matrice avec les cours assignés
        for i, c in enumerate(cours):
            # Find assigned room
            salle_assignee = None
            
            # Vérifier si on a une assignation personnalisée pour ce cours
            course_id = c.get('id')
            if custom_room_assignments and course_id:
                try:
                    # S'assurer que custom_room_assignments est un dictionnaire
                    if isinstance(custom_room_assignments, str):
                        import json
                        custom_room_assignments = json.loads(custom_room_assignments)
                    
                    if isinstance(custom_room_assignments, dict):
                        assigned_room_name = custom_room_assignments.get(str(course_id))
                        if assigned_room_name:
                            # Trouver la clé de la salle à partir du nom
                            for room_key, room_data in salles.items():
                                if room_data.get('nom') == assigned_room_name or room_key == assigned_room_name:
                                    salle_assignee = room_key
                                    print(f"🎯 Assignation personnalisée: cours {course_id} -> salle {room_key}")
                                    break
                except Exception as e:
                    print(f"⚠️ Erreur traitement assignations personnalisées: {e}")
                    # En cas d'erreur, continuer avec l'assignation normale
            
            # Si pas d'assignation personnalisée, utiliser l'assignation du solver
            if not salle_assignee:
                for s in salle_list:
                    if (i, s) in x and solver.Value(x[(i, s)]) == 1:
                        salle_assignee = s
                        break
            
            if salle_assignee:
                # Parse course time
                horaire_debut = c.get('horaire', '8:00')
                start_time = h_to_min(horaire_debut)
                duree = c.get('duree', 110)  # Durée par défaut
                
                # Find start and end indices in horaires
                idx_start = 0
                idx_end = 1
                
                # Find the closest start time
                for idx, h in enumerate(horaires):
                    if h <= start_time:
                        idx_start = idx
                    if h >= start_time + duree and idx_end == 1:
                        idx_end = idx
                        break
                else:
                    # Si on n'a pas trouvé de fin, utiliser la fin du tableau
                    idx_end = len(horaires)
                
                # S'assurer que idx_end > idx_start
                if idx_end <= idx_start:
                    idx_end = idx_start + 1
                
                merge_len = idx_end - idx_start
                
                # Extraire les équipements de salle demandés depuis selected_materials
                besoins_equipements = extract_material_needs(c.get('selected_materials', ''))
                
                # Build course content for Planning_Techniciens (équipements de salle + titre TP)
                besoins_txt_techniciens = []
                
                # Ajouter seulement les équipements de salle qui sont réellement demandés
                if besoins_equipements.get('eviers', 0) > 0:
                    besoins_txt_techniciens.append("Éviers")
                if besoins_equipements.get('hotte', 0) > 0:
                    besoins_txt_techniciens.append("Hotte")
                if besoins_equipements.get('bancs_optiques', 0) > 0:
                    besoins_txt_techniciens.append("Bancs optiques")
                if besoins_equipements.get('obscurite_totale', 0) > 0:
                    besoins_txt_techniciens.append("Obscurité totale")
                if besoins_equipements.get('becs_electriques', 0) > 0:
                    besoins_txt_techniciens.append("Becs électriques")
                if besoins_equipements.get('support_filtration', 0) > 0:
                    besoins_txt_techniciens.append("Support de filtration")
                if besoins_equipements.get('imprimante', 0) > 0:
                    besoins_txt_techniciens.append("Imprimante")
                
                # Ajouter le titre du TP (request_name) s'il existe
                titre_tp = c.get('request_name', '')
                if titre_tp and titre_tp.strip():
                    besoins_txt_techniciens.append(titre_tp.strip())
                
                rest_lines_techniciens = [c.get('niveau', '')]
                if besoins_txt_techniciens:
                    rest_lines_techniciens.append(", ".join(besoins_txt_techniciens))
                content_techniciens = _teacher_rich_text(c.get('enseignant', ''), rest_lines_techniciens)

                # Build course content for Affichage (seulement enseignant + niveau)
                content_affichage = _teacher_rich_text(c.get('enseignant', ''), [c.get('niveau', '')])
                
                # Fill matrix avec vérifications d'indices
                try:
                    salle_idx = salle_list.index(salle_assignee)
                except ValueError:
                    continue  # Salle non trouvée, ignorer
                
                # Vérifier que les indices sont dans les limites
                if (idx_start >= 0 and idx_start < len(cell_matrix) and 
                    salle_idx >= 0 and salle_idx < len(salle_list)):
                    
                    # Remplir seulement la première cellule pour éviter les conflits avec les cellules fusionnées
                    if idx_start < len(cell_matrix) and salle_idx < len(cell_matrix[idx_start]):
                        cell_matrix[idx_start][salle_idx]["content_techniciens"] = content_techniciens
                        cell_matrix[idx_start][salle_idx]["content_affichage"] = content_affichage
                        cell_matrix[idx_start][salle_idx]["content"] = content_techniciens  # Par défaut, pour compatibilité
                        cell_matrix[idx_start][salle_idx]["matiere"] = c.get('matiere', 'mixte')
                        
                        # Marquer les autres cellules comme "occupées" mais sans contenu
                        for idx in range(idx_start + 1, min(idx_end, len(cell_matrix))):
                            if idx < len(cell_matrix) and salle_idx < len(cell_matrix[idx]):
                                cell_matrix[idx][salle_idx]["content_techniciens"] = ""  # Cellule occupée mais vide
                                cell_matrix[idx][salle_idx]["content_affichage"] = ""
                                cell_matrix[idx][salle_idx]["content"] = ""  # Cellule occupée mais vide
                                cell_matrix[idx][salle_idx]["matiere"] = c.get('matiere', 'mixte')
                    
                    # Marquer pour fusion seulement si valide
                    if merge_len > 1 and idx_start < len(cell_matrix):
                        cell_matrix[idx_start][salle_idx]["merge"] = True
                        cell_matrix[idx_start][salle_idx]["merge_len"] = merge_len
        
                # Remplissage du tableau avec fusion verticale
        # D'abord, faire toutes les fusions
        merged_ranges = set()
        for idx_h, h in enumerate(horaires):
            for idx_s, s in enumerate(salle_list):
                cell = cell_matrix[idx_h][idx_s]
                excel_row = 3 + idx_h
                excel_col = 2 + idx_s
                
                # Fusion des cellules - seulement pour les cellules de début avec contenu
                if cell["merge"] and cell["content"] and cell["merge_len"] > 1:
                    end_row = excel_row + cell["merge_len"] - 1
                    merge_key = (excel_row, excel_col, end_row)
                    
                    # Vérifier que la fusion est valide et pas déjà faite
                    if (end_row > excel_row and 
                        end_row <= 3 + len(horaires) - 1 and 
                        merge_key not in merged_ranges):
                        
                        try:
                            ws1.merge_cells(
                                start_row=excel_row, start_column=excel_col,
                                end_row=end_row, end_column=excel_col
                            )
                            merged_ranges.add(merge_key)
                        except Exception as e:
                            print(f"Erreur fusion: start_row={excel_row}, end_row={end_row}, col={excel_col}: {e}")
        
        # Ensuite, remplir le contenu
        for idx_h, h in enumerate(horaires):
            row_border = _row_border(idx_h)
            # Horaires (colonne 1)
            ws1.cell(row=3+idx_h, column=1, value=horaires_str[idx_h])
            ws1.cell(row=3+idx_h, column=1).alignment = Alignment(horizontal="center", vertical="center")
            ws1.cell(row=3+idx_h, column=1).border = row_border

            # Salles
            for idx_s, s in enumerate(salle_list):
                cell = cell_matrix[idx_h][idx_s]
                excel_row = 3 + idx_h
                excel_col = 2 + idx_s

                # Vérifier si cette cellule est une cellule fusionnée secondaire
                is_merged_secondary = False
                for merge_key in merged_ranges:
                    start_row, start_col, end_row = merge_key
                    if (start_col == excel_col and
                        start_row < excel_row <= end_row):
                        is_merged_secondary = True
                        break

                # Appliquer les bordures à toutes les cellules (y compris fusionnées)
                try:
                    ws1.cell(row=excel_row, column=excel_col).border = row_border
                except Exception:
                    pass
                
                # Remplir seulement les cellules principales (pas les secondaires)
                if not is_merged_secondary:
                    try:
                        # Utiliser le contenu techniciens pour la feuille Planning_Techniciens
                        ws1.cell(row=excel_row, column=excel_col, value=cell["content_techniciens"])
                        ws1.cell(row=excel_row, column=excel_col).alignment = Alignment(
                            horizontal="center", vertical="center", wrap_text=True
                        )
                        
                        # Couleurs seulement si contenu
                        if cell["content"]:
                            matiere = cell["matiere"].lower() if cell["matiere"] else ""
                            if "chimie" in matiere:
                                ws1.cell(row=excel_row, column=excel_col).fill = PatternFill(
                                    start_color="B2F2E9", end_color="B2F2E9", fill_type="solid"
                                )
                            elif "physique" in matiere:
                                ws1.cell(row=excel_row, column=excel_col).fill = PatternFill(
                                    start_color="FFD580", end_color="FFD580", fill_type="solid"
                                )
                            else:
                                ws1.cell(row=excel_row, column=excel_col).fill = PatternFill(
                                    start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
                                )
                    except Exception as e:
                        # Ignorer les erreurs sur les cellules fusionnées
                        print(f"Ignoring cell write error at ({excel_row}, {excel_col}): {e}")
                        pass        # Hauteur des lignes
        for idx_h in range(len(horaires)):
            ws1.row_dimensions[3 + idx_h].height = 25

        # Feuille 2: Affichage simplifié
        ws2 = wb.create_sheet(title="Affichage")

        # Même structure mais contenu simplifié
        ws2.cell(row=1, column=1, value=jour_semaine)
        ws2.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws2.column_dimensions['A'].width = 12
        ws2.row_dimensions[1].height = 40
        ws2.row_dimensions[2].height = 40

        # Répliquer les en-têtes
        _ecrire_entetes_de_groupes(ws2, groupes_salles, salle_list, header_font, border)

        # Noms des salles (ligne 2)
        for idx_s, s in enumerate(salle_list):
            col_letter = ws2.cell(row=2, column=2+idx_s).column_letter
            cell = ws2.cell(row=2, column=2+idx_s, value=s)
            ws2.column_dimensions[col_letter].width = 15
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        
        # Date (ligne 2, colonne 1)
        ws2.cell(row=2, column=1, value=formatted_date)
        ws2.cell(row=2, column=1).font = bold_font
        ws2.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=2, column=1).border = border
        
        # Version simplifiée pour l'affichage (seulement enseignants)
        cell_matrix2 = [[{"content": "", "merge": False, "merge_len": 1, "matiere": ""} for _ in salle_list] for _ in horaires]
        
        for i, c in enumerate(cours):
            # Find assigned room
            salle_assignee = None
            
            # Vérifier si on a une assignation personnalisée pour ce cours (même logique que Planning_Techniciens)
            course_id = c.get('id')
            if custom_room_assignments and course_id:
                try:
                    # S'assurer que custom_room_assignments est un dictionnaire
                    if isinstance(custom_room_assignments, str):
                        import json
                        custom_room_assignments = json.loads(custom_room_assignments)
                    
                    if isinstance(custom_room_assignments, dict):
                        assigned_room_name = custom_room_assignments.get(str(course_id))
                        if assigned_room_name:
                            # Trouver la clé de la salle à partir du nom
                            for room_key, room_data in salles.items():
                                if room_data.get('nom') == assigned_room_name or room_key == assigned_room_name:
                                    salle_assignee = room_key
                                    print(f"🎯 [Affichage] Assignation personnalisée: cours {course_id} -> salle {room_key}")
                                    break
                except Exception as e:
                    print(f"⚠️ [Affichage] Erreur traitement assignations personnalisées: {e}")
                    # En cas d'erreur, continuer avec l'assignation normale
            
            # Si pas d'assignation personnalisée, utiliser l'assignation du solver
            if not salle_assignee:
                for s in salle_list:
                    if (i, s) in x and solver.Value(x[(i, s)]) == 1:
                        salle_assignee = s
                        break
            
            if salle_assignee:
                horaire_debut = c.get('horaire', '8:00')
                start_time = h_to_min(horaire_debut)
                duree = c.get('duree', 110)
                
                # Find start and end indices in horaires (même logique)
                idx_start = 0
                idx_end = 1
                
                for idx, h in enumerate(horaires):
                    if h <= start_time:
                        idx_start = idx
                    if h >= start_time + duree and idx_end == 1:
                        idx_end = idx
                        break
                else:
                    idx_end = len(horaires)
                
                if idx_end <= idx_start:
                    idx_end = idx_start + 1
                
                merge_len = idx_end - idx_start
                
                # Build course content for Affichage (seulement enseignant + niveau)
                content_affichage = _teacher_rich_text(c.get('enseignant', ''), [c.get('niveau', '')])

                # Fill matrix avec vérifications d'indices
                try:
                    salle_idx = salle_list.index(salle_assignee)
                except ValueError:
                    continue  # Salle non trouvée, ignorer
                
                # Vérifier que les indices sont dans les limites
                if (idx_start >= 0 and idx_start < len(cell_matrix2) and 
                    salle_idx >= 0 and salle_idx < len(salle_list)):
                    
                    # Remplir seulement la première cellule pour éviter les conflits avec les cellules fusionnées
                    if idx_start < len(cell_matrix2) and salle_idx < len(cell_matrix2[idx_start]):
                        # Utiliser le contenu d'affichage (seulement enseignant + niveau)
                        cell_matrix2[idx_start][salle_idx]["content"] = content_affichage
                        cell_matrix2[idx_start][salle_idx]["matiere"] = c.get('matiere', 'mixte')
                        
                        # Marquer les autres cellules comme "occupées" mais sans contenu
                        for idx in range(idx_start + 1, min(idx_end, len(cell_matrix2))):
                            if idx < len(cell_matrix2) and salle_idx < len(cell_matrix2[idx]):
                                cell_matrix2[idx][salle_idx]["content"] = ""  # Cellule occupée mais vide
                                cell_matrix2[idx][salle_idx]["matiere"] = c.get('matiere', 'mixte')
                    
                    # Marquer pour fusion seulement si valide
                    if merge_len > 1 and idx_start < len(cell_matrix2):
                        cell_matrix2[idx_start][salle_idx]["merge"] = True
                        cell_matrix2[idx_start][salle_idx]["merge_len"] = merge_len
        
        # Remplissage feuille 2 - même logique séparée
        # D'abord, faire toutes les fusions
        merged_ranges2 = set()
        for idx_h, h in enumerate(horaires):
            for idx_s, s in enumerate(salle_list):
                cell = cell_matrix2[idx_h][idx_s]
                excel_row = 3 + idx_h
                excel_col = 2 + idx_s
                
                # Fusion des cellules - seulement pour les cellules de début avec contenu
                if cell["merge"] and cell["content"] and cell["merge_len"] > 1:
                    end_row = excel_row + cell["merge_len"] - 1
                    merge_key = (excel_row, excel_col, end_row)
                    
                    # Vérifier que la fusion est valide et pas déjà faite
                    if (end_row > excel_row and 
                        end_row <= 3 + len(horaires) - 1 and 
                        merge_key not in merged_ranges2):
                        
                        try:
                            ws2.merge_cells(
                                start_row=excel_row, start_column=excel_col,
                                end_row=end_row, end_column=excel_col
                            )
                            merged_ranges2.add(merge_key)
                        except Exception as e:
                            print(f"Erreur fusion feuille 2: start_row={excel_row}, end_row={end_row}, col={excel_col}: {e}")
        
        # Ensuite, remplir le contenu
        for idx_h, h in enumerate(horaires):
            row_border = _row_border(idx_h)
            ws2.cell(row=3+idx_h, column=1, value=horaires_str[idx_h])
            ws2.cell(row=3+idx_h, column=1).alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(row=3+idx_h, column=1).border = row_border

            for idx_s, s in enumerate(salle_list):
                cell = cell_matrix2[idx_h][idx_s]
                excel_row = 3 + idx_h
                excel_col = 2 + idx_s

                # Vérifier si cette cellule est une cellule fusionnée secondaire
                is_merged_secondary = False
                for merge_key in merged_ranges2:
                    start_row, start_col, end_row = merge_key
                    if (start_col == excel_col and
                        start_row < excel_row <= end_row):
                        is_merged_secondary = True
                        break

                # Appliquer les bordures à toutes les cellules (y compris fusionnées)
                try:
                    ws2.cell(row=excel_row, column=excel_col).border = row_border
                except Exception:
                    pass
                
                # Remplir seulement les cellules principales (pas les secondaires)
                if not is_merged_secondary:
                    try:
                        ws2.cell(row=excel_row, column=excel_col, value=cell["content"])
                        ws2.cell(row=excel_row, column=excel_col).alignment = Alignment(
                            horizontal="center", vertical="center", wrap_text=True
                        )
                        
                        if cell["content"]:
                            ws2.cell(row=excel_row, column=excel_col).fill = PatternFill(
                                start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
                            )
                    except Exception as e:
                        # Ignorer les erreurs sur les cellules fusionnées
                        print(f"Ignoring cell write error feuille 2 at ({excel_row}, {excel_col}): {e}")
                        pass
        
        # Hauteur des lignes feuille 2
        for idx_h in range(len(horaires)):
            ws2.row_dimensions[3 + idx_h].height = 25

        # Mise en page A4 paysage pour impression
        _apply_print_setup(ws1, len(salle_list), len(horaires))
        _apply_print_setup(ws2, len(salle_list), len(horaires))

        # Ajout des cours non assignés en bas
        if unassigned_courses:
            # Ligne vide
            ws1.cell(row=3 + len(horaires) + 1, column=1, value="COURS NON ASSIGNÉS")
            ws1.cell(row=3 + len(horaires) + 1, column=1).font = Font(bold=True, color='FF0000', size=14)
            
            row_start = 3 + len(horaires) + 2
            for i, course_info in enumerate(unassigned_courses):
                ws1.cell(row=row_start + i, column=1, value=course_info)
                ws1.cell(row=row_start + i, column=1).font = Font(color='FF0000')
        
        # Sauvegarde du fichier
        from datetime import datetime
        filename = f"planning_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        wb.save(filename)
        
        assigned_count = len(cours) - len(unassigned_courses)
        print(f"Planning généré: {filename} ({assigned_count}/{len(cours)} cours assignés)")
        return True, filename
        
    except Exception as e:
        return False, f"Erreur lors de la génération Excel: {str(e)}"


def generer_excel_from_saved_planning(planning_data, date_str):
    """
    Génère l'Excel directement depuis le planning sauvegardé en base (table `plannings`),
    sans relancer OR-Tools. Respecte les modifications de salle ET de créneau réalisées
    via glisser-déposer dans l'éditeur.

    planning_data : dict tel que stocké par /api/save-planning
        { courses: [{room, time, duration, teacher, level, subject,
                     request_name, eviers, hotte, bancs_optiques,
                     obscurite_totale, becs_electriques, support_filtration,
                     imprimante, ...}], ... }
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
        from openpyxl.worksheet.page import PageMargins
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from datetime import datetime

        JOURS_FR = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche']

        def _row_border(idx_h):
            """Bordure horizontale alternée (pleine/pointillés), avec séparateurs
            renforcés aux lignes de pause réglées dans /admin/establishment."""
            is_major = horaires_str[idx_h] in separateurs
            next_is_major = idx_h + 1 < len(horaires_str) and horaires_str[idx_h + 1] in separateurs
            top_style = 'thick' if is_major else ('medium' if idx_h % 2 == 0 else 'dashed')
            bottom_style = 'thick' if next_is_major else ('medium' if idx_h % 2 == 0 else 'dashed')
            return Border(left=Side(style='thick'), right=Side(style='thick'),
                          top=Side(style=top_style), bottom=Side(style=bottom_style))

        def _teacher_rich_text(teacher, rest_lines, highlight_lines=None):
            """Nom du professeur en gras, 1.5x plus grand que le reste du contenu.
            highlight_lines (effectif, ordinateurs) reprennent la même mise en forme
            que le nom (gras, grand), une par ligne, comme dans l'éditeur de planning.
            Un élément de highlight_lines peut être un tuple (texte, couleur_argb)
            pour surligner un cas particulier (ex : dépassement du nombre d'ordinateurs)."""
            rest = "\n".join(l for l in rest_lines if l)
            runs = [TextBlock(InlineFont(rFont="Calibri", b=True, sz=17), teacher or "")]
            if rest:
                runs.append(TextBlock(InlineFont(rFont="Calibri", sz=11), "\n" + rest))
            for hl in (highlight_lines or []):
                if not hl:
                    continue
                color = None
                if isinstance(hl, tuple):
                    hl, color = hl
                runs.append(TextBlock(InlineFont(rFont="Calibri", b=True, sz=17, color=color), "\n" + str(hl)))
            return CellRichText(*runs)

        def _apply_print_setup(ws, n_rooms, last_row):
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3, header=0, footer=0)
            last_col_letter = ws.cell(row=2, column=1 + n_rooms).column_letter
            ws.print_area = f"A1:{last_col_letter}{last_row}"

        courses = planning_data.get('courses', [])
        room_assignments = planning_data.get('room_assignments', {})

        # Rafraîchir depuis la base les infos qui peuvent avoir changé depuis l'enregistrement du
        # planning (noms courts, observations labo, effectifs), comme dans l'éditeur de planning.
        rafraichir_cours(courses)

        # Demandes "Test au labo" (pas de salle attribuée, donc absentes de `courses`)
        # et observations internes du planning : à afficher sous la grille.
        lab_tests = database.get_lab_test_requests(date_str)
        observations = database.get_planning_observations(date_str)

        wb = Workbook()

        layout_horaires = site_settings.get_time_layout()
        horaires_str = layout_horaires['grid_rows']
        separateurs = set(layout_horaires['separators'])
        horaires = [h_to_min(h) for h in horaires_str]

        ordered_rooms, groupes_salles = get_room_layout()

        # Salles présentes dans le planning sauvegardé (course.room ou room_assignments)
        rooms_in_planning = set()
        for c in courses:
            r = c.get('room', '') or room_assignments.get(str(c.get('id', '')), '')
            if r:
                rooms_in_planning.add(r)
        all_known_rooms = set(ordered_rooms)
        salle_list = [r for r in ordered_rooms if r in rooms_in_planning or r in all_known_rooms]
        # Ajouter les salles inconnues
        for r in rooms_in_planning:
            if r not in salle_list:
                salle_list.append(r)

        bold_font   = Font(bold=True, size=16)
        header_font = Font(bold=True, size=36)
        border = Border(
            left=Side(style='thick'), right=Side(style='thick'),
            top=Side(style='thick'), bottom=Side(style='thick')
        )

        if date_str:
            try:
                planning_date  = datetime.strptime(date_str, '%Y-%m-%d')
                formatted_date = planning_date.strftime('%d/%m')
                jour_semaine   = JOURS_FR[planning_date.weekday()]
            except ValueError:
                formatted_date = date_str
                jour_semaine   = date_str
        else:
            planning_date  = datetime.now()
            formatted_date = planning_date.strftime('%d/%m')
            jour_semaine   = JOURS_FR[planning_date.weekday()]

        def _build_sheet(wb, title, use_techniciens_content):
            ws = wb.active if title == "Planning_Techniciens" else wb.create_sheet(title=title)
            ws.title = title

            ws.cell(row=1, column=1, value=jour_semaine).alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions['A'].width = 12
            ws.row_dimensions[1].height = 40
            ws.row_dimensions[2].height = 40

            _ecrire_entetes_de_groupes(ws, groupes_salles, salle_list, header_font, border)

            for idx_s, s in enumerate(salle_list):
                col_letter = ws.cell(row=2, column=2 + idx_s).column_letter
                cell = ws.cell(row=2, column=2 + idx_s, value=s)
                ws.column_dimensions[col_letter].width = 15
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

            ws.cell(row=2, column=1, value=formatted_date)
            ws.cell(row=2, column=1).font = bold_font
            ws.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=2, column=1).border = border

            cell_matrix = [
                [{"content": "", "merge": False, "merge_len": 1, "matiere": ""} for _ in salle_list]
                for _ in horaires
            ]

            for c in courses:
                room = c.get('room', '') or room_assignments.get(str(c.get('id', '')), '')
                if not room or room not in salle_list:
                    continue

                # Créneau depuis la position dans la grille (mis à jour par drag)
                time_str   = c.get('time', '9h00') or '9h00'
                start_time = h_to_min(time_str)
                duree      = int(c.get('duration', 85) or 85)

                idx_start = 0
                idx_end   = 1
                for idx, h in enumerate(horaires):
                    if h <= start_time:
                        idx_start = idx
                    if h >= start_time + duree and idx_end == 1:
                        idx_end = idx
                        break
                else:
                    idx_end = len(horaires)
                if idx_end <= idx_start:
                    idx_end = idx_start + 1

                merge_len = idx_end - idx_start
                salle_idx = salle_list.index(room)
                matiere   = c.get('subject', 'mixte') or 'mixte'
                level     = c.get('level', '')

                if use_techniciens_content:
                    # Même contenu que la carte de cours de l'éditeur de planning :
                    # nom court, niveau court, observations labo, effectif (2nde classe
                    # entière) et nombre d'ordinateurs.
                    teacher = c.get('teacher_short_name') or c.get('teacher', '')
                    level_display = c.get('level_short_name') or level
                    rest_lines = [level_display]
                    labo_obs = (c.get('labo_observations') or '').strip()
                    if labo_obs:
                        rest_lines.append(labo_obs)
                    highlight_lines = []
                    if level == '2nd Classe' and c.get('students', 0):
                        highlight_lines.append(c.get('students'))
                    if c.get('ordinateurs', 0):
                        ordi_text = f"💻 {c.get('ordinateurs')}"
                        if c.get('ordinateurs_overflow'):
                            highlight_lines.append((ordi_text, "FFDC3545"))
                        else:
                            highlight_lines.append(ordi_text)
                    content = _teacher_rich_text(teacher, rest_lines, highlight_lines)
                else:
                    teacher = c.get('teacher', '')
                    content = _teacher_rich_text(teacher, [level])

                if idx_start < len(cell_matrix) and salle_idx < len(cell_matrix[idx_start]):
                    cell_matrix[idx_start][salle_idx]["content"]    = content
                    cell_matrix[idx_start][salle_idx]["matiere"]    = matiere
                    cell_matrix[idx_start][salle_idx]["merge"]      = merge_len > 1
                    cell_matrix[idx_start][salle_idx]["merge_len"]  = merge_len
                    for idx in range(idx_start + 1, min(idx_end, len(cell_matrix))):
                        cell_matrix[idx][salle_idx]["content"] = ""
                        cell_matrix[idx][salle_idx]["matiere"] = matiere

            # Fusions
            merged_ranges = set()
            for idx_h in range(len(horaires)):
                for idx_s in range(len(salle_list)):
                    cell = cell_matrix[idx_h][idx_s]
                    excel_row = 3 + idx_h
                    excel_col = 2 + idx_s
                    if cell["merge"] and cell["content"] and cell["merge_len"] > 1:
                        end_row   = excel_row + cell["merge_len"] - 1
                        merge_key = (excel_row, excel_col, end_row)
                        if end_row > excel_row and end_row <= 3 + len(horaires) - 1 and merge_key not in merged_ranges:
                            try:
                                ws.merge_cells(start_row=excel_row, start_column=excel_col,
                                               end_row=end_row,   end_column=excel_col)
                                merged_ranges.add(merge_key)
                            except Exception:
                                pass

            # Contenu
            for idx_h, h in enumerate(horaires):
                row_border = _row_border(idx_h)
                ws.cell(row=3 + idx_h, column=1, value=horaires_str[idx_h])
                ws.cell(row=3 + idx_h, column=1).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=3 + idx_h, column=1).border    = row_border

                for idx_s in range(len(salle_list)):
                    cell      = cell_matrix[idx_h][idx_s]
                    excel_row = 3 + idx_h
                    excel_col = 2 + idx_s

                    is_secondary = any(
                        sr < excel_row <= er and sc == excel_col
                        for (sr, sc, er) in merged_ranges
                    )
                    try:
                        ws.cell(row=excel_row, column=excel_col).border = row_border
                    except Exception:
                        pass

                    if not is_secondary:
                        try:
                            ws.cell(row=excel_row, column=excel_col, value=cell["content"])
                            ws.cell(row=excel_row, column=excel_col).alignment = Alignment(
                                horizontal="center", vertical="center", wrap_text=True)
                            if cell["content"]:
                                mat = (cell["matiere"] or "").lower()
                                color = "B2F2E9" if "chimie" in mat else "FFD580" if "physique" in mat else "D3D3D3"
                                ws.cell(row=excel_row, column=excel_col).fill = PatternFill(
                                    start_color=color, end_color=color, fill_type="solid")
                        except Exception:
                            pass

            for idx_h in range(len(horaires)):
                ws.row_dimensions[3 + idx_h].height = 25

            last_row = 2 + len(horaires)

            # Tests au labo (sans salle) et observations internes : uniquement sur la
            # feuille technicien·ne·s, sous la grille (comme dans l'éditeur de planning).
            if use_techniciens_content:
                last_col = 1 + len(salle_list)
                footer_row = last_row + 2

                if lab_tests:
                    lab_lines = []
                    for lt in lab_tests:
                        piece = lt.get('teacher_name', '') or ''
                        if lt.get('horaire'):
                            piece = f"{lt['horaire']} — {piece}"
                        if lt.get('class_name'):
                            piece += f" ({lt['class_name']})"
                        if lt.get('request_name'):
                            piece += f" — {lt['request_name']}"
                        lab_lines.append(piece)
                    cell = ws.cell(row=footer_row, column=1,
                                    value="🧪 Tests au labo (sans salle attribuée) : " + " ; ".join(lab_lines))
                    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=last_col)
                    cell.font = Font(bold=True, size=12)
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws.row_dimensions[footer_row].height = 30
                    footer_row += 1

                if observations and observations.strip():
                    cell = ws.cell(row=footer_row, column=1, value="Observations : " + observations.strip())
                    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=last_col)
                    cell.font = Font(size=12, italic=True)
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws.row_dimensions[footer_row].height = 40
                    footer_row += 1

                last_row = max(last_row, footer_row - 1)

            _apply_print_setup(ws, len(salle_list), last_row)

        _build_sheet(wb, "Planning_Techniciens", use_techniciens_content=True)
        _build_sheet(wb, "Affichage",             use_techniciens_content=False)

        import tempfile, os
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        wb.save(tmp.name)
        return True, tmp.name

    except Exception as e:
        import traceback
        traceback.print_exc()
        return False, f"Erreur génération Excel depuis planning sauvegardé: {e}"
